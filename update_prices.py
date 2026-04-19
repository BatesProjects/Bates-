#!/usr/bin/env python3
"""
Bunnings Price Updater
Reads Price Book.xlsx, scrapes live prices from Bunnings, and updates the spreadsheet.

Usage:
  python3 update_prices.py "Price Book.xlsx"        # Run on all rows
  python3 update_prices.py --test "Price Book.xlsx" # Test on first 5 URLs only
"""

import argparse
import json
import logging
import re
import shutil
import sys
import time
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import PatternFill

# Playwright is the preferred scraping engine (uses a real browser).
# If not installed, the script falls back to plain HTTP requests.
try:
    from playwright.sync_api import sync_playwright
    PLAYWRIGHT_AVAILABLE = True
except ImportError:
    PLAYWRIGHT_AVAILABLE = False

# ─── Configuration ─────────────────────────────────────────────────────────────

QUOTE_SHEET_NAME = "Quote"
LOG_SHEET_NAME   = "Price Change Log"

URL_COLUMN    = 2   # Column B — Bunnings product URLs
PRICE_COLUMN  = 5   # Column E — unit price
DATE_COLUMN   = 7   # Column G — date last checked
DATA_START_ROW = 2  # Row 1 is the header row

DELAY_SECONDS = 2   # Pause between requests (increase if still getting blocks)
TEST_LIMIT    = 5   # URLs processed in --test mode

# ─── Fill colours ──────────────────────────────────────────────────────────────
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
RED_FILL    = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
NO_FILL     = PatternFill(fill_type="none")

# ─── Browser headers for the requests fallback ─────────────────────────────────
BROWSER_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept":          "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-AU,en;q=0.9",
    "Accept-Encoding": "gzip, deflate, br",
    "Connection":      "keep-alive",
    "Referer":         "https://www.bunnings.com.au/",
}


# ─── Helpers ───────────────────────────────────────────────────────────────────

def backup_spreadsheet(path: Path) -> Path:
    timestamp  = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = Path.home() / "Desktop" / "Price Book Backups"
    backup_dir.mkdir(exist_ok=True)
    backup = backup_dir / f"{path.stem}_backup_{timestamp}{path.suffix}"
    shutil.copy2(path, backup)
    logging.info("Backup saved: %s", backup)
    return backup


def save_workbook(wb: openpyxl.Workbook, spreadsheet_path: Path) -> Path:
    """Save workbook; if iCloud blocks the write, save to Desktop instead."""
    try:
        wb.save(spreadsheet_path)
        logging.info("Saved: %s", spreadsheet_path)
        return spreadsheet_path
    except PermissionError:
        fallback = Path.home() / "Desktop" / spreadsheet_path.name
        wb.save(fallback)
        logging.warning(
            "iCloud blocked saving to the original location.\n"
            "  Updated file saved to your Desktop: %s\n"
            "  Please drag it back into your Price Book folder to replace the old one.",
            fallback,
        )
        return fallback


def highlight_row(ws, row_idx: int, fill: PatternFill) -> None:
    for col in range(1, ws.max_column + 1):
        ws.cell(row=row_idx, column=col).fill = fill


def get_or_create_log_sheet(wb: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    if LOG_SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(LOG_SHEET_NAME)
        ws.append([
            "Row", "Item Description", "URL",
            "Old Price ($)", "New Price ($)", "Change ($)",
            "Date Checked", "Status",
        ])
    return wb[LOG_SHEET_NAME]


def _extract_price_from_html(html: str) -> Optional[float]:
    """Parse price out of a fully-rendered HTML string."""
    soup = BeautifulSoup(html, "html.parser")

    # Method 1: JSON-LD structured data
    for tag in soup.find_all("script", type="application/ld+json"):
        try:
            data  = json.loads(tag.string or "")
            items = data if isinstance(data, list) else [data]
            for item in items:
                if item.get("@type") == "Product":
                    offers = item.get("offers") or {}
                    if isinstance(offers, list):
                        offers = offers[0]
                    price_val = offers.get("price")
                    if price_val is not None:
                        return float(str(price_val).replace(",", ""))
        except (json.JSONDecodeError, ValueError, TypeError):
            continue

    # Method 2: microdata / common Bunnings CSS selectors
    for selector in [
        "[itemprop='price']",
        "[data-locator='price-display']",
        "[data-testid='price-display']",
        "[data-testid='price']",
        ".price__value",
        ".pdp-price__value",
        "span[class*='Price']",
    ]:
        el = soup.select_one(selector)
        if el:
            content = el.get("content") or el.get_text(strip=True)
            m = re.search(r"[\d,]+\.\d{2}", content.replace("$", ""))
            if m:
                try:
                    return float(m.group().replace(",", ""))
                except ValueError:
                    pass

    # Method 3: JavaScript variables in the page source
    for tag in soup.find_all("script"):
        text = tag.string or ""
        if not any(k in text for k in ("sellPrice", "currentPrice", '"price"')):
            continue
        m = re.search(
            r'"(?:sellPrice|currentPrice|price)"\s*:\s*([\d]+(?:\.\d+)?)', text
        )
        if m:
            try:
                candidate = float(m.group(1))
                if 0.01 <= candidate <= 100_000:
                    return candidate
            except ValueError:
                pass

    return None


# ─── Scraping ──────────────────────────────────────────────────────────────────

def _scrape_with_playwright(url: str) -> Optional[float]:
    """Use a real headless Chrome browser — bypasses Cloudflare / JS rendering."""
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            context = browser.new_context(
                user_agent=(
                    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/120.0.0.0 Safari/537.36"
                ),
                locale="en-AU",
                timezone_id="Australia/Sydney",
            )
            page = context.new_page()
            page.goto(url, wait_until="domcontentloaded", timeout=30_000)
            page.wait_for_timeout(3_000)  # let JS finish rendering prices

            # Try the microdata price attribute directly via the browser DOM
            price_el = page.locator("[itemprop='price']").first
            if price_el.count() > 0:
                content = price_el.get_attribute("content")
                if content:
                    try:
                        return float(content.replace(",", ""))
                    except ValueError:
                        pass

            # Try other common selectors
            for selector in [
                "[data-testid='price-display']",
                "[data-locator='price-display']",
                ".price__value",
                ".pdp-price__value",
                "span[class*='Price']",
            ]:
                els = page.locator(selector)
                if els.count() > 0:
                    text = els.first.inner_text()
                    m = re.search(r"[\d,]+\.\d{2}", text.replace("$", ""))
                    if m:
                        try:
                            return float(m.group().replace(",", ""))
                        except ValueError:
                            pass

            # Fall back to full HTML parse
            return _extract_price_from_html(page.content())

        except Exception as exc:
            logging.warning("  Playwright error: %s", exc)
            return None
        finally:
            browser.close()


def _scrape_with_requests(url: str, session: requests.Session) -> Optional[float]:
    """Plain HTTP fallback — may be blocked by Cloudflare."""
    try:
        resp = session.get(url, headers=BROWSER_HEADERS, timeout=15)
    except requests.RequestException as exc:
        logging.warning("  Network error: %s", exc)
        return None

    if resp.status_code in (403, 429, 503):
        logging.warning("  Blocked (HTTP %s) — Cloudflare may be active.", resp.status_code)
        return None
    if not resp.ok:
        logging.warning("  HTTP %s", resp.status_code)
        return None

    return _extract_price_from_html(resp.text)


def scrape_bunnings_price(url: str, session: requests.Session) -> Optional[float]:
    if PLAYWRIGHT_AVAILABLE:
        return _scrape_with_playwright(url)
    return _scrape_with_requests(url, session)


# ─── Main logic ────────────────────────────────────────────────────────────────

def run(spreadsheet_path: Path, test_mode: bool = False) -> None:
    if not spreadsheet_path.exists():
        logging.error("Spreadsheet not found: %s", spreadsheet_path)
        sys.exit(1)

    logging.info("=" * 65)
    logging.info("Bunnings Price Updater — %s", datetime.now().strftime("%Y-%m-%d %H:%M"))
    if test_mode:
        logging.info("TEST MODE  (first %d URLs only)", TEST_LIMIT)
    logging.info("Scraping engine: %s", "Playwright (browser)" if PLAYWRIGHT_AVAILABLE else "requests (fallback)")
    logging.info("Spreadsheet: %s", spreadsheet_path)
    logging.info("=" * 65)

    backup_spreadsheet(spreadsheet_path)

    wb = openpyxl.load_workbook(spreadsheet_path)

    if QUOTE_SHEET_NAME not in wb.sheetnames:
        logging.error(
            "Sheet '%s' not found. Sheets in this file: %s",
            QUOTE_SHEET_NAME, wb.sheetnames,
        )
        sys.exit(1)

    ws     = wb[QUOTE_SHEET_NAME]
    log_ws = get_or_create_log_sheet(wb)
    today  = date.today().isoformat()

    session     = requests.Session()
    url_count   = 0
    success     = 0
    changed     = 0
    failed      = 0

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        url = ws.cell(row=row_idx, column=URL_COLUMN).value
        if not url or not isinstance(url, str):
            continue
        url = url.strip()
        if not url.lower().startswith("http"):
            continue

        url_count += 1

        old_price_raw = ws.cell(row=row_idx, column=PRICE_COLUMN).value
        try:
            old_price: Optional[float] = float(old_price_raw) if old_price_raw is not None else None
        except (ValueError, TypeError):
            old_price = None

        desc = ws.cell(row=row_idx, column=1).value or ""
        logging.info("Row %d | %s", row_idx, url[:80])

        if not PLAYWRIGHT_AVAILABLE:
            time.sleep(DELAY_SECONDS)

        new_price = scrape_bunnings_price(url, session)

        if new_price is None:
            failed += 1
            highlight_row(ws, row_idx, RED_FILL)
            log_ws.append([
                row_idx, str(desc), url,
                old_price, "FAILED", "N/A",
                today, "FAILED — could not retrieve price",
            ])
            logging.warning("  => FAILED  Row %d highlighted RED, existing price kept.", row_idx)
        else:
            success += 1
            ws.cell(row=row_idx, column=PRICE_COLUMN).value = new_price
            ws.cell(row=row_idx, column=DATE_COLUMN).value  = today

            price_changed = old_price is not None and abs(new_price - old_price) > 0.005
            change_amount = round(new_price - old_price, 2) if old_price is not None else "N/A"

            if price_changed:
                changed += 1
                highlight_row(ws, row_idx, YELLOW_FILL)
                status = f"CHANGED (was ${old_price:.2f})"
            else:
                highlight_row(ws, row_idx, NO_FILL)
                status = "UNCHANGED"

            log_ws.append([
                row_idx, str(desc), url,
                old_price, new_price, change_amount,
                today, status,
            ])
            logging.info(
                "  => $%.2f  (was %s)  [%s]",
                new_price,
                f"${old_price:.2f}" if old_price is not None else "unknown",
                status,
            )

        if test_mode and url_count >= TEST_LIMIT:
            logging.info("Test mode complete — processed %d URLs.", url_count)
            break

    save_workbook(wb, spreadsheet_path)

    logging.info("=" * 65)
    logging.info(
        "Done.  Checked: %d  |  Updated: %d  |  Changed: %d  |  Failed: %d",
        url_count, success, changed, failed,
    )
    logging.info("=" * 65)


def main() -> None:
    parser = argparse.ArgumentParser(description="Update Bunnings prices in Price Book.xlsx")
    parser.add_argument(
        "spreadsheet",
        nargs="?",
        default="Price Book.xlsx",
        help="Path to the spreadsheet (default: 'Price Book.xlsx' in current folder)",
    )
    parser.add_argument(
        "--test",
        action="store_true",
        help=f"Test mode: only process the first {TEST_LIMIT} URLs",
    )
    args = parser.parse_args()

    spreadsheet_path = Path(args.spreadsheet)
    if not spreadsheet_path.is_absolute() and not spreadsheet_path.exists():
        candidate = Path(__file__).parent / spreadsheet_path
        if candidate.exists():
            spreadsheet_path = candidate

    if not PLAYWRIGHT_AVAILABLE:
        print(
            "\nNOTE: Playwright is not installed — using basic HTTP scraping.\n"
            "If prices fail to load, install Playwright for better results:\n"
            "  pip3 install playwright\n"
            "  python3 -m playwright install chromium\n"
        )

    log_file = Path.home() / "Desktop" / "price_updater.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
        handlers=[
            logging.StreamHandler(sys.stdout),
            logging.FileHandler(log_file, encoding="utf-8"),
        ],
    )

    run(spreadsheet_path, test_mode=args.test)


if __name__ == "__main__":
    main()
